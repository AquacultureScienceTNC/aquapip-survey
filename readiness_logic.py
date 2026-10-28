"""
AquaPIP - "How close are you?" readiness logic (no Streamlit; used by the page and the PDFs).

Reads the 'Readiness' sheet of MEL_Indicator_Reference.xlsx by header name, so rows,
questions and wording can be edited in Excel without touching code.

Scoring (agreed Oct 2026):
  * Only the core measurement that produces an indicator's metric earns points. Doing it is
    worth BASE_POINTS (60). Indicators the Framework splits into parts (e.g. community
    composition: gear + mobile animals + seabed) are scaled by the share of parts covered.
  * The other 40 points come from up to three follow-up questions (where / how often /
    records or detail) scored against the Framework minimum for that indicator.
  * Until a question is answered the score is a range (low = no credit, high = full credit).
  * Context measurements are listed but never scored.
"""

import os
import openpyxl

SHEET = "Readiness"
BASE_POINTS = 60

BANDS = [  # key, name, range label, lower bound on the rounded midpoint
    ("near", "Nearly there", "85\u2013100%", 85),
    ("part", "Partway there", "50\u201384%", 50),
    ("first", "First steps", "1\u201349%", 1),
    ("new", "New for you", "Nothing you do feeds these yet", 0),
]
BAND_NAME = {k: n for k, n, _r, _m in BANDS}


# --------------------------------------------------------------------------- #
# Loading
# --------------------------------------------------------------------------- #
def _s(v):
    return "" if v is None else str(v).strip()


def _semis(v):
    return [x.strip() for x in _s(v).split(";") if x.strip()]


def _options(v):
    """'Only inside the farm=0.5; Inside the farm and at a spot outside it=1' -> [(label, credit)]"""
    out = []
    for item in _semis(v):
        if "=" not in item:
            continue
        label, credit = item.rsplit("=", 1)
        try:
            c = float(credit)
        except ValueError:
            continue
        out.append((label.strip(), max(0.0, min(1.0, c))))
    return out


def _float(v, default):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def load(ref_path):
    """{(sector, code): recipe dict}. Empty dict if the sheet is missing (the page then
    explains that readiness isn't set up yet, and the rest of the app runs unchanged)."""
    if not ref_path or not os.path.exists(ref_path):
        return {}
    wb = openpyxl.load_workbook(ref_path, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        return {}
    it = wb[SHEET].iter_rows(values_only=True)
    try:
        header = [_s(h).lower() for h in next(it)]
    except StopIteration:
        return {}
    col = {h: i for i, h in enumerate(header) if h}

    def cell(row, name):
        i = col.get(name.lower())
        return row[i] if i is not None and i < len(row) else None

    out = {}
    for row in it:
        if not row:
            continue
        sector, code = _s(cell(row, "Sector")), _s(cell(row, "Indicator Code"))
        if not sector or not code:
            continue
        parts = []
        for n in (1, 2, 3):
            label = _s(cell(row, "Part %d label" % n))
            if label:
                parts.append({"label": label, "methods": _semis(cell(row, "Part %d methods" % n)),
                              "hint": _s(cell(row, "Part %d hint" % n))})
        questions = []
        for n in (1, 2, 3):
            text, opts = _s(cell(row, "Q%d text" % n)), _options(cell(row, "Q%d options" % n))
            if text and opts:
                questions.append({"key": "q%d" % n, "text": text, "options": opts,
                                  "weight": max(0.0, _float(cell(row, "Q%d weight" % n), 1.0)),
                                  "gap": _s(cell(row, "Q%d gap" % n))})
        out[(sector, code)] = {
            "sector": sector, "code": code, "name": _s(cell(row, "Indicator Name")) or code,
            "picture": _s(cell(row, "Picture")), "tagline": _s(cell(row, "Tagline")),
            "summary": _s(cell(row, "Plain-language summary")), "caption": _s(cell(row, "Picture caption")),
            "parts": parts, "context": _semis(cell(row, "Context methods")), "questions": questions,
            "first": _s(cell(row, "First step")),
            "first_short": _s(cell(row, "First step (short)")) or _s(cell(row, "First step")),
            "note": _s(cell(row, "Note when in place")),
        }
    return out


def recipe_for(recipes, sectors, code):
    """First recipe for `code` among the farm's sectors (IMTA farms span several)."""
    for sec in sectors:
        r = recipes.get((sec, code))
        if r:
            return r
    return None


def sector_codes(recipes, sectors):
    seen, out = set(), []
    for sec in sectors:
        for (s, code) in recipes:
            if s == sec and code not in seen:
                seen.add(code); out.append(code)
    return out


# --------------------------------------------------------------------------- #
# Scoring
# --------------------------------------------------------------------------- #
def score(recipe, have, answers=None, base=BASE_POINTS):
    """have: set of screener labels the farmer ticked. answers: {question key: option index}."""
    a = answers or {}
    parts = recipe["parts"]
    covered = sum(1 for p in parts if any(m in have for m in p["methods"]))
    m = covered / len(parts) if parts else 0.0
    rest = 100.0 - base
    totw = sum(q["weight"] for q in recipe["questions"])
    got = 0.0 if totw else rest
    open_pts, open_n = 0.0, 0
    for q in recipe["questions"]:
        pts = rest * q["weight"] / totw if totw else 0.0
        k = a.get(q["key"])
        if k is None or not (0 <= int(k) < len(q["options"])):
            open_pts += pts; open_n += 1
        else:
            got += pts * q["options"][int(k)][1]
    lo = m * (base + got)
    hi = m * (base + got + open_pts)
    return {"M": m, "covered": covered, "n_parts": len(parts), "lo": lo, "hi": hi,
            "mid": (lo + hi) / 2.0, "open": open_pts, "open_n": open_n if m > 0 else 0,
            "got": got, "rest": rest, "totw": totw}


def band(sc):
    mid = round(sc["mid"])
    for key, _n, _r, lower in BANDS:
        if mid >= lower:
            return key
    return "new"


def score_text(sc):
    lo, hi = round(sc["lo"]), round(sc["hi"])
    if sc["M"] <= 0:
        return "0%"
    return "%d\u2013%d%%" % (lo, hi) if hi > lo else "%d%%" % lo


def core_methods(recipe, have):
    out = []
    for p in recipe["parts"]:
        for m in p["methods"]:
            if m in have and m not in out:
                out.append(m)
    return out


def context_methods(recipe, have):
    core = set(core_methods(recipe, have))
    return [m for m in recipe["context"] if m in have and m not in core]


def todos(recipe, sc, have, answers=None, base=BASE_POINTS):
    """[(text, gain text)] for missing parts and answers that earn less than full credit."""
    a = answers or {}
    out = []
    if sc["M"] <= 0:
        return out
    for p in recipe["parts"]:
        if not any(m in have for m in p["methods"]):
            gain = round((1.0 / sc["n_parts"]) * (base + sc["got"] + sc["open"] / 2.0))
            txt = "Add " + p["label"][0].lower() + p["label"][1:] + (": " + p["hint"] if p["hint"] else "") + "."
            out.append((txt, "about +%d%%" % gain if gain > 0 else ""))
    for q in recipe["questions"]:
        k = a.get(q["key"])
        if k is None or not (0 <= int(k) < len(q["options"])):
            continue
        credit = q["options"][int(k)][1]
        if credit < 1:
            pts = sc["rest"] * q["weight"] / sc["totw"] if sc["totw"] else 0
            gain = round(sc["M"] * pts * (1 - credit))
            out.append((q["gap"] or q["text"], "+%d%%" % gain if gain > 0 else ""))
    return out


def question_points(recipe, base=BASE_POINTS):
    totw = sum(q["weight"] for q in recipe["questions"])
    rest = 100.0 - base
    return {q["key"]: (round(rest * q["weight"] / totw) if totw else 0) for q in recipe["questions"]}


def rank(entries):
    """entries: list of (code, recipe, score). Closest first; framework order breaks ties."""
    return sorted(enumerate(entries), key=lambda t: (-t[1][2]["mid"], -t[1][2]["lo"], t[0]))


# --------------------------------------------------------------------------- #
# Page model
# --------------------------------------------------------------------------- #
def active_codes(base_codes, removed, added):
    out = [c for c in base_codes if c not in removed]
    out += [c for c in added if c not in out]
    return out


def suggestions(recipes, sectors, active, removed, have, answers=None):
    """Indicators not on the list whose core measurement the farmer already covers (at least one part)."""
    answers = answers or {}
    out = []
    for code in sector_codes(recipes, sectors):
        if code in active or code in removed:
            continue
        r = recipe_for(recipes, sectors, code)
        sc = score(r, have, answers.get(code))
        if sc["covered"] > 0:
            out.append((code, r, sc))
    return out


def readout(recipes, sectors, codes, have, answers=None):
    """Ranked rows for the printable readout and page 1 of the summary PDF."""
    answers = answers or {}
    entries = []
    for code in codes:
        r = recipe_for(recipes, sectors, code)
        if r:
            entries.append((code, r, score(r, have, answers.get(code))))
    rows = []
    for _i, (code, r, sc) in rank(entries):
        core = core_methods(r, have)
        tds = todos(r, sc, have, answers.get(code))
        if sc["M"] <= 0:
            nl, nt = "First step:", r["first_short"]
        elif tds:
            nl, nt = "Next:", " ".join(t.rstrip(".") + (" (%s)" % g if g else "") + "." for t, g in tds)
        elif sc["open_n"]:
            nl, nt = "Next:", "%d question%s not answered yet, so this shows a range." % (sc["open_n"], "" if sc["open_n"] == 1 else "s")
        else:
            nl, nt = "Next:", "Nothing needed; you meet the Framework\u2019s minimum."
        rows.append({
            "code": code, "name": r["name"], "tagline": r["tagline"], "summary": r["summary"],
            "caption": r["caption"], "picture": r["picture"], "score": score_text(sc),
            "lo": sc["lo"], "hi": sc["hi"], "band": BAND_NAME[band(sc)],
            "in_place": (", ".join(core) + (" (%d of %d parts)" % (sc["covered"], sc["n_parts"]) if 0 < sc["M"] < 1 else "")) if core else "Nothing yet",
            "next_label": nl, "next_text": nt,
            "open_note": ("%d question%s not answered yet, so this shows a range." % (sc["open_n"], "" if sc["open_n"] == 1 else "s"))
                         if sc["open_n"] and tds else "",
        })
    return rows


def dial(recipes, sectors, codes, have, answers=None):
    """(answered, total) follow-up questions across indicators the farm has started."""
    answers = answers or {}
    done = total = 0
    for code in codes:
        r = recipe_for(recipes, sectors, code)
        if not r:
            continue
        sc = score(r, have, answers.get(code))
        if sc["M"] > 0:
            total += len(r["questions"])
            done += len(r["questions"]) - sc["open_n"]
    return done, total
